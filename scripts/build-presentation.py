#!/usr/bin/env python3
"""Build the defense slides as editable PPTX and matching PDF.

Optional documentation dependencies: python-pptx, reportlab, Pillow.
Run from any directory; outputs are written to docs/.
"""
from pathlib import Path
from html import escape
import os

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.lib.styles import ParagraphStyle
from reportlab.platypus import Paragraph

ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / 'docs'
W, H = 13.333333, 7.5
BG, INK, GREEN, MUTED, PALE = 'F7F8F5', '243C33', '117665', '607260', 'EAF3EB'
font_candidates = [
    (Path(os.environ.get('WINDIR', 'C:/Windows')) / 'Fonts/arial.ttf',
     Path(os.environ.get('WINDIR', 'C:/Windows')) / 'Fonts/arialbd.ttf'),
    (Path('/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf'),
     Path('/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf')),
]
font_regular, font_bold = next((a, b) for a, b in font_candidates if a.exists() and b.exists())
pdfmetrics.registerFont(TTFont('Body', str(font_regular)))
pdfmetrics.registerFont(TTFont('Bold', str(font_bold)))
prs = Presentation()
prs.slide_width, prs.slide_height = Inches(W), Inches(H)
prs.core_properties.title = 'Траты — учёт личных расходов'
prs.core_properties.subject = 'Защита семестрового проекта по Go'
prs.core_properties.author = 'El1syum'
pdf = canvas.Canvas(str(DOCS / 'presentation.pdf'), pagesize=(W*72, H*72))
pdf.setTitle(prs.core_properties.title)
pdf.setAuthor('El1syum')
slide = None


def rect(x, y, w, h, color):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor.from_string(color)
    shape.line.fill.background()
    pdf.setFillColorRGB(*(int(color[i:i+2], 16)/255 for i in (0, 2, 4)))
    pdf.rect(x*72, (H-y-h)*72, w*72, h*72, fill=1, stroke=0)


def text(value, x, y, w, h, size=20, color=INK, bold=False):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.text = value
    for p in tf.paragraphs:
        p.font.name = 'Arial'
        p.font.size = Pt(size)
        p.font.bold = bold
        p.font.color.rgb = RGBColor.from_string(color)
        p.space_after = Pt(0)
        p.line_spacing = 1.2
    style = ParagraphStyle('text', fontName='Bold' if bold else 'Body', fontSize=size,
                           leading=size*1.2, textColor='#'+color)
    paragraph = Paragraph(escape(value).replace('\n', '<br/>'), style)
    _, height = paragraph.wrap(w*72, h*72)
    if height > h*72 + 1:
        raise ValueError(f'Text overflow on slide {len(prs.slides)}: {value[:60]}')
    paragraph.drawOn(pdf, x*72, (H-y)*72-height)


def picture(name, x, y, w, h):
    path = DOCS / 'screenshots' / name
    with Image.open(path) as im:
        ratio = im.width / im.height
    pw, ph = w, w/ratio
    if ph > h:
        ph, pw = h, h*ratio
    x += (w-pw)/2
    y += (h-ph)/2
    slide.shapes.add_picture(str(path), Inches(x), Inches(y), width=Inches(pw), height=Inches(ph))
    pdf.drawImage(str(path), x*72, (H-y-ph)*72, pw*72, ph*72)


def page(kicker, title, notes):
    global slide
    if slide is not None:
        pdf.showPage()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.notes_slide.notes_text_frame.text = notes
    rect(0, 0, W, H, BG)
    rect(.5, .5, .08, .2, GREEN)
    text(kicker.upper(), .72, .47, 11.8, .25, 10, GREEN, True)
    text(title, .7, 1.0, 12, .7, 32, INK, True)
    rect(.7, 6.99, 11.9, .012, 'DDE5DC')
    text('ТРАТЫ  /  СЕМЕСТРОВЫЙ ПРОЕКТ НА GO', .7, 7.15, 10, .2, 9, MUTED)
    text(f'{len(prs.slides):02d} / 07', 11.9, 7.15, .75, .2, 9, MUTED)


page('01 / Задача', 'Личные расходы — в понятной системе',
     '0:00–0:35. Представьте приложение и задачу. Все восемь этапов объединены в один серверный Go-проект. '
     'Пользователи ведут собственные расходы, анализируют их и планируют бюджет. Назовите опубликованный адрес.')
text('траты.', .7, 2.1, 4.7, 1.0, 54, GREEN, True)
text('Записывать. Понимать.\nПланировать.', .7, 3.3, 4.5, 1.1, 26, INK, True)
text('8 этапов + 4 дополнительных задания\nGo · SQLite · серверный HTML', .7, 4.7, 4.5, 1.1, 18, MUTED)
text('integradar.org', .7, 6.0, 4.6, .4, 23, GREEN, True)
picture('expenses.png', 5.35, 2.0, 7.3, 4.6)

page('02 / Пользовательский сценарий', 'От первой записи до общей картины',
     '0:35–1:25. Покажите регистрацию и вход, затем готовую историю. Добавьте небольшую трату, '
     'измените описание, откройте подтверждение удаления и отмените. Примените категорию и даты. '
     'Откройте статистику и CSV. Все данные принадлежат текущему пользователю.')
for n, value in enumerate(['Регистрация и вход', 'Добавление и изменение трат', 'Категория + период', 'Графики и CSV']):
    text(f'0{n+1}', .7, 2.15+n*.95, .7, .5, 22, GREEN, True)
    text(value, 1.5, 2.16+n*.95, 3.4, .8, 20)
picture('stats.png', 5.15, 2.0, 7.5, 4.75)

page('03 / Архитектура', 'У каждого слоя — своя задача',
     '1:25–2:20. Проследите запрос: HTTPS/Caddy → net/http → middleware → handler → repository → SQLite. '
     'Объясните ResponseWriter и коды ответа. HTML формируют шаблоны с общим layout; после успешного POST — 303. '
     'Временный срез этапа 2 заменён SQLite на этапе 3. SQL находится только в репозиториях.')
labels = ['Браузер\nHTML / fetch', 'Go HTTP\nMiddleware', 'Handlers\nВалидация', 'Repository\nSQLite']
for i, label in enumerate(labels):
    x = .7+i*3.03
    rect(x, 2.15, 2.75, 1.3, PALE)
    text(label, x+.2, 2.42, 2.35, .85, 21, GREEN, True)
    if i < 3:
        text('→', x+2.78, 2.52, .3, .5, 20, MUTED)
text('HTML', .7, 4.05, 3.5, .4, 21, GREEN, True)
text('html/template + layout\nЭкранирование данных\nPost/Redirect/Get', .7, 4.65, 3.5, 1.5, 20)
text('ДАННЫЕ', 4.75, 4.05, 3.5, .4, 21, GREEN, True)
text('database/sql + миграции\nПараметры SQL через ?\nСвязи и индексы', 4.75, 4.65, 3.5, 1.5, 20)
text('ИНФРАСТРУКТУРА', 8.8, 4.05, 3.9, .4, 21, GREEN, True)
text('Docker + постоянная БД\n.env и graceful shutdown\nCaddy + HTTPS', 8.8, 4.65, 3.8, 1.5, 20)

page('04 / Технические решения', 'Точность данных и разделение доступа',
     '2:20–3:20. Деньги — int64 в копейках, 0.10 + 0.20 = 0.30. JOIN исключает N+1, '
     'SUM/GROUP BY считают статистику в БД. Пароль — bcrypt, cookie содержит случайный токен, в базе HMAC. '
     'Middleware проверяет срок сессии и передаёт пользователя через context. Чужой ID даёт 404. '
     'Расскажите о CSRF, экранировании HTML и отзыве сессии при выходе.')
items = [
    ('Целые копейки', 'Точный разбор сумм;\nокругление не накапливается.'),
    ('JOIN + SUM / GROUP BY', 'Категории без N+1;\nагрегации выполняет SQLite.'),
    ('Сессии и bcrypt', 'Срок действия, отзыв, HttpOnly;\nuser_id только из контекста.'),
    ('Проверка каждого запроса', 'Серверная валидация, CSRF;\nчужие записи недоступны.'),
]
for i, (head, body) in enumerate(items):
    x, y = .7+(i%2)*6.08, 2.1+(i//2)*2.3
    rect(x, y, 5.8, 2.0, 'FFFFFF')
    text(head, x+.25, y+.23, 5.3, .5, 24, GREEN, True)
    text(body, x+.25, y+.9, 5.3, 1.0, 20)

page('05 / Дополнительные задания', 'Все четыре возможности работают',
     '3:20–4:35. Покажите бюджет с превышением. Создайте месячную подписку на сегодня — запись появится сразу. '
     'Пауза пропускает даты; завершённое расписание можно отредактировать. Транзакция и уникальная пара '
     'расписание/дата предотвращают дубли даже после рестарта; удаление сохраняет историю. '
     'Покажите тёмную тему. Docker — реальный способ запуска опубликованного сайта.')
items = [
    ('01 / Бюджеты', 'Лимит категории на месяц.\nОстаток и предупреждение.'),
    ('02 / Регулярные траты', 'День, неделя, месяц, год.\nПауза, окончание, без дублей.'),
    ('03 / Тёмная тема', 'Светлая, тёмная, системная.\nВыбор сохраняется в браузере.'),
    ('04 / Docker', 'Non-root, volume, healthcheck.\nПроверка чистой БД в CI.'),
]
for i, (head, body) in enumerate(items):
    x, y = .7+(i%2)*6.08, 2.1+(i//2)*2.3
    rect(x, y, 5.8, 2.0, PALE)
    text(head, x+.25, y+.23, 5.3, .5, 24, GREEN, True)
    text(body, x+.25, y+.9, 5.3, 1.0, 20)

page('06 / Проверка', 'Проверены сценарии и границы',
     '4:35–5:40. Покажите Go-тесты, go vet и GitHub Actions с race detector, сборкой Docker и HTTP smoke. '
     'Smoke создаёт двух пользователей и проходит весь путь плюс дополнительные функции. '
     'Отдельно проверяются пустые списки, неверные суммы и даты, 404/500, XSS, CSRF, CSV-формулы. '
     'Для расписаний проверены 31-е число, 29 февраля, конкурентные запуски и откат SQL-ошибки. '
     'Миграция старой базы сохраняет расходы. Покажите реальный JSON-лог из verification.md.')
text('АВТОМАТИЗАЦИЯ', .7, 2.15, 5.6, .4, 18, GREEN, True)
text('go test -race -cover ./...\ngo vet ./...\nDocker build + HTTP smoke', .7, 2.85, 5.6, 1.65, 23)
text('Два аккаунта: CRUD, фильтры, API,\nCSV, бюджеты, расписания, тема.', .7, 5.05, 5.6, 1.2, 20, MUTED)
rect(6.85, 2.1, 5.75, 4.55, 'FFFFFF')
text('ГРАНИЧНЫЕ СЛУЧАИ', 7.15, 2.4, 5.1, .4, 18, GREEN, True)
text('Пустой список и неверный ввод\nЧужой ID, истёкшая сессия\n31-е число и 29 февраля\nПовторный запуск без дублей\nОбновление существующей БД', 7.15, 3.2, 5.0, 3.0, 22)

page('07 / Результат', 'Готово к запуску и демонстрации',
     '5:40–6:40. Откройте публичный сайт и покажите ключевые экраны. Назовите репозиторий и README. '
     'Кратко объясните запуск через go run . и docker compose up --build -d. '
     'Все инструкции и проверка каждого пункта ТЗ находятся в docs. Завершите вопросами преподавателя.')
text('integradar.org', .7, 2.2, 4.6, .6, 30, GREEN, True)
text('github.com/El1syum/tksu_pl', .7, 3.1, 4.6, .55, 18)
text('README: запуск и база проекта\nАудит каждого пункта ТЗ\nПрезентация и сценарий защиты\nТесты и скриншоты', .7, 4.0, 4.6, 2.0, 19, MUTED)
picture('dark-theme.png', 5.45, 2.0, 7.2, 4.65)

prs.save(DOCS / 'presentation.pptx')
pdf.save()
assert len(prs.slides) == 7
print('Created docs/presentation.pptx and docs/presentation.pdf (7 slides).')
